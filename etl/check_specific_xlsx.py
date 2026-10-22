import openpyxl
import os

parts_files = [
    r"c:\Users\techb\Desktop\electrol-portal-NEW-main\electrol-portal-NEW-main\uptaed voter list\100_1791587428.xlsx",
    r"c:\Users\techb\Desktop\electrol-portal-NEW-main\electrol-portal-NEW-main\uptaed voter list\101_1791587543.xlsx",
    r"c:\Users\techb\Desktop\electrol-portal-NEW-main\electrol-portal-NEW-main\uptaed voter list\102_1791587543.xlsx",
    r"c:\Users\techb\Desktop\electrol-portal-NEW-main\electrol-portal-NEW-main\uptaed voter list\103_1791587543.xlsx",
]

for pf in parts_files:
    if os.path.exists(pf):
        wb = openpyxl.load_workbook(pf, data_only=True)
        ws = wb.active
        print(f"\n--- Checking {os.path.basename(pf)} ({ws.max_row} rows) ---")
        for r_idx in range(1, ws.max_row + 1):
            row_vals = [ws.cell(row=r_idx, column=c).value for c in range(1, ws.max_column + 1)]
            row_str = " ".join([str(v) for v in row_vals if v is not None])
            if "4252136" in row_str or "YNT4252136" in row_str:
                print(f"  FOUND in {os.path.basename(pf)} row {r_idx}:", row_vals)
            elif "4252" in row_str and "YNT" in row_str:
                print(f"  Near match in row {r_idx}:", row_vals[:5])
        wb.close()
