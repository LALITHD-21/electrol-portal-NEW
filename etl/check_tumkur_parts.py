import os
import glob
import openpyxl
import fitz

target_epic = "YNT4252136"
target_num = "4252136"

print("--- Checking Tumkur parts in 'uptaed voter list' ---")
base_dir = r"c:\Users\techb\Desktop\electrol-portal-NEW-main\electrol-portal-NEW-main"

tumkur_files = []
for p in range(98, 106):
    tumkur_files.extend(glob.glob(os.path.join(base_dir, "uptaed voter list", f"{p}_*.xlsx")))
    tumkur_files.extend(glob.glob(os.path.join(base_dir, "new file", f"{p}_*.xlsx")))

print(f"Found {len(tumkur_files)} Tumkur xlsx files:")
for tf in tumkur_files:
    print(f"Scanning {os.path.basename(tf)}...")
    try:
        wb = openpyxl.load_workbook(tf, read_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                r_str = " ".join([str(c) for c in row if c is not None])
                if target_num in r_str or target_epic in r_str:
                    print(f"  >>> FOUND in {os.path.basename(tf)} row {r_idx} <<<")
                    print("  Data:", row)
        wb.close()
    except Exception as e:
        print(f"  Error: {e}")

print("\n--- Checking Tumkur PDFs in VOTERS PDF ---")
tumkur_pdfs = []
for p in range(98, 106):
    tumkur_pdfs.extend(glob.glob(os.path.join(base_dir, "VOTERS PDF", f"{p}_*.pdf")))
    tumkur_pdfs.extend(glob.glob(os.path.join(base_dir, "updates voter list pdf", f"{p}_*.pdf")))

print(f"Found {len(tumkur_pdfs)} Tumkur PDF files:")
for pf in tumkur_pdfs:
    print(f"Scanning PDF {os.path.basename(pf)}...")
    try:
        doc = fitz.open(pf)
        for p_idx in range(len(doc)):
            txt = doc[p_idx].get_text()
            if target_num in txt or target_epic in txt:
                print(f"  >>> FOUND in {os.path.basename(pf)} page {p_idx+1} <<<")
                for line in txt.split('\n'):
                    if target_num in line or "YNT" in line:
                        print("    ", line)
        doc.close()
    except Exception as e:
        print(f"  Error reading PDF: {e}")

