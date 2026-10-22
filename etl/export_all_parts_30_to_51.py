import sqlite3
import re
import os
import shutil
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TALUK_MAP = {
    30: ("Molakalmuru", "Molakalmuru"),
    31: ("Molakalmuru", "Molakalmuru"),
    32: ("Challakere", "Challakere"),
    33: ("Challakere", "Challakere"),
    34: ("Challakere", "Challakere"),
    35: ("Challakere", "Challakere"),
    36: ("Chitradurga", "Chitradurga"),
    37: ("Chitradurga", "Chitradurga"),
    38: ("Chitradurga", "Chitradurga"),
    39: ("Chitradurga", "Chitradurga"),
    40: ("Hiriyur", "Hiriyur"),
    41: ("Hiriyur", "Hiriyur"),
    42: ("Hiriyur", "Hiriyur"),
    43: ("Hiriyur", "Hiriyur"),
    44: ("Hosadurga", "Hosadurga"),
    45: ("Hosadurga", "Hosadurga"),
    46: ("Hosadurga", "Hosadurga"),
    47: ("Hosadurga", "Hosadurga"),
    48: ("Holalkere", "Holalkere"),
    49: ("Holalkere", "Holalkere"),
    50: ("Holalkere", "Holalkere"),
    51: ("Holalkere", "Holalkere"),
}

def clean_name(val):
    if not val:
        return ""
    s = str(val).strip()
    if s.lower() in ('null', 'none', '—', '-', '.'):
        return ""
    # Strip OCR symbols
    s = re.sub(r'[\$#@!~%^*_+=\[\]{};:\"\'\<\>?/\\|]+', '', s).strip()
    # Separate lowercase letters followed directly by uppercase letters (e.g. "RojaY" -> "Roja Y")
    s = re.sub(r'([a-z])([A-Z])', r'\1 \2', s)
    # Collapse multiple whitespace
    s = re.sub(r'\s+', ' ', s).strip()
    return s

def clean_phone(val):
    if not val:
        return ""
    digits = re.sub(r'\D', '', str(val))
    if len(digits) == 10 and digits[0] in '6789':
        if len(set(digits)) > 1:
            return digits
    return ""

def clean_address(val):
    if not val:
        return ""
    s = str(val).strip()
    if s.lower() in ('null', 'none', '—', '-'):
        return ""
    s = s.replace('\\n', ', ').replace('\n', ', ').replace('\r', '')
    s = re.sub(r'[,;\s]+,', ',', s)
    s = re.sub(r',\s*,+', ', ', s)
    s = re.sub(r'\s+', ' ', s).strip(' ,')
    return s

def clean_text(val):
    if not val:
        return ""
    s = str(val).strip()
    if s.lower() in ('null', 'none', '—', '-'):
        return ""
    return re.sub(r'\s+', ' ', s)

def export_part_file(part_num, out_dir, conn):
    c = conn.cursor()
    c.execute("""
        SELECT 
            id,
            serial_number,
            epic_number,
            name,
            relative_name,
            sex,
            age,
            whatsapp_mob,
            address,
            part_number,
            polling_station_name,
            polling_address,
            qualification,
            occupation
        FROM electors
        WHERE CAST(part_number AS INTEGER) = ?
        ORDER BY 
            CASE WHEN serial_number IS NOT NULL AND serial_number != '' THEN 0 ELSE 1 END,
            CAST(serial_number AS INTEGER),
            id
    """, (part_num,))
    
    rows = c.fetchall()
    total_electors = len(rows)
    taluk_name, ac_name = TALUK_MAP.get(part_num, ("Chitradurga", "Chitradurga"))
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Part {part_num} Voters"
    
    headers = [
        "Sl No",
        "Roll Serial No",
        "EPIC Number",
        "Elector Name",
        "Relative Name",
        "Gender",
        "Age",
        "Mobile Number",
        "Address",
        "Part Number",
        "Polling Station Name",
        "Polling Address",
        "Taluk",
        "District",
        "Assembly Constituency",
        "Qualification",
        "Occupation"
    ]
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_side = Side(style='thin', color='D1D5DB')
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    regular_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = cell_border
        
    phone_count = 0
    center_cols = {1, 2, 3, 6, 7, 8, 10}
    
    for idx, r in enumerate(rows, start=1):
        row_id, r_serial, r_epic, r_name, r_rel, r_sex, r_age, r_phone, r_addr, r_part, r_station, r_paddr, r_qual, r_occ = r
        
        c_roll_serial = str(r_serial).strip() if (r_serial is not None and str(r_serial).strip() != '') else None
        c_epic = str(r_epic or '').strip().upper()
        c_name = clean_name(r_name)
        c_rel = clean_name(r_rel)
        c_sex = str(r_sex or '').strip().upper()
        if c_sex in ('M', 'MALE'): c_sex = 'M'
        elif c_sex in ('F', 'FEMALE'): c_sex = 'F'
        elif c_sex in ('T', 'THIRD', 'OTHER'): c_sex = 'T'
        
        c_age = int(r_age) if (r_age is not None and str(r_age).isdigit()) else None
        c_phone = clean_phone(r_phone)
        if c_phone:
            phone_count += 1
            phone_val = c_phone
        else:
            phone_val = None  # None becomes empty blank cell in openpyxl
            
        c_addr = clean_address(r_addr)
        c_part = str(part_num)
        c_station = clean_text(r_station)
        c_paddr = clean_text(r_paddr)
        c_qual = clean_text(r_qual)
        c_occ = clean_text(r_occ)
        
        row_data = [
            idx,
            c_roll_serial,
            c_epic,
            c_name,
            c_rel,
            c_sex,
            c_age,
            phone_val,
            c_addr,
            c_part,
            c_station,
            c_paddr,
            taluk_name,
            "Chitradurga",
            ac_name,
            c_qual,
            c_occ
        ]
        
        ws.append(row_data)
        current_row = idx + 1
        ws.row_dimensions[current_row].height = 20
        
        for c_idx in range(1, len(headers) + 1):
            c_cell = ws.cell(row=current_row, column=c_idx)
            c_cell.font = regular_font
            c_cell.border = cell_border
            c_cell.alignment = center_align if c_idx in center_cols else left_align
            # Format phone as text so Excel doesn't trim leading zeros or use scientific notation
            if c_idx == 8 and phone_val:
                c_cell.number_format = '@'

    # Freeze header row
    ws.freeze_panes = "A2"
    
    # Auto-filter
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col_letter}{total_electors + 1}"
    
    min_widths = {
        1: 8,   # Sl No
        2: 14,  # Roll Serial No
        3: 16,  # EPIC Number
        4: 26,  # Elector Name
        5: 26,  # Relative Name
        6: 9,   # Gender
        7: 8,   # Age
        8: 16,  # Mobile Number
        9: 45,  # Address
        10: 12, # Part Number
        11: 35, # Polling Station Name
        12: 40, # Polling Address
        13: 16, # Taluk
        14: 15, # District
        15: 22, # Assembly Constituency
        16: 18, # Qualification
        17: 20  # Occupation
    }
    
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min_widths.get(col_idx, 15)
        
    out_file = os.path.join(out_dir, f"Part_{part_num}_Voters.xlsx")
    wb.save(out_file)
    wb.close()
    return total_electors, phone_count, out_file

def main():
    start_time = time.time()
    db_path = 'c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/web/electors.db'
    out_dir_1 = 'c:/Users/techb/Desktop/electrol-portal-NEW-main/exports/Parts_30_to_51'
    out_dir_2 = 'c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/exports/Parts_30_to_51'
    
    os.makedirs(out_dir_1, exist_ok=True)
    os.makedirs(out_dir_2, exist_ok=True)
    
    conn = sqlite3.connect(db_path)
    
    summary = []
    print("=========================================================================")
    print("STARTING EXCEL GENERATION FOR PART NUMBERS 30 TO 51")
    print("=========================================================================")
    
    total_all_voters = 0
    total_all_phones = 0
    
    for part in range(30, 52):
        t0 = time.time()
        voters, phones, fpath = export_part_file(part, out_dir_1, conn)
        # Copy to second export directory
        fpath_2 = os.path.join(out_dir_2, os.path.basename(fpath))
        shutil.copy2(fpath, fpath_2)
        
        elapsed = time.time() - t0
        total_all_voters += voters
        total_all_phones += phones
        taluk = TALUK_MAP.get(part, ("Chitradurga", "Chitradurga"))[0]
        
        summary.append({
            "part": part,
            "filename": os.path.basename(fpath),
            "taluk": taluk,
            "voters": voters,
            "phones": phones,
            "phone_pct": round((phones / voters * 100), 1) if voters > 0 else 0,
            "time_sec": round(elapsed, 2)
        })
        print(f"[{part}/51] Part {part:>2} ({taluk:<12}): {voters:>5} voters | {phones:>4} with phones ({phones/voters*100:4.1f}%) | {elapsed:.2f}s")
        
    conn.close()
    total_elapsed = time.time() - start_time
    
    print("=========================================================================")
    print(f"EXPORT COMPLETED SUCCESSFULLY IN {total_elapsed:.1f}s")
    print(f"Total Parts Exported: {len(summary)} (Parts 30 to 51)")
    print(f"Total Electors Exported: {total_all_voters:,}")
    print(f"Total Electors with Verified Phone Numbers: {total_all_phones:,} ({total_all_phones/total_all_voters*100:.1f}%)")
    print(f"Export Directory 1: {out_dir_1}")
    print(f"Export Directory 2: {out_dir_2}")
    print("=========================================================================")

if __name__ == '__main__':
    main()
