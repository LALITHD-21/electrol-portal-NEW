import os
import shutil
import zipfile
import sqlite3
import re
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

start_time = time.time()

# 1. Target Single Folder on Desktop
target_desktop_folder = r"C:\Users\techb\Desktop\Parts_30_to_51_Voters"
os.makedirs(target_desktop_folder, exist_ok=True)

source_dir = r"C:\Users\techb\Desktop\electrol-portal-NEW-main\exports\Parts_30_to_51"

print(f"Copying all 22 individual Excel files to {target_desktop_folder}...")
copied_files = []
for part in range(30, 52):
    fname = f"Part_{part}_Voters.xlsx"
    src = os.path.join(source_dir, fname)
    dst = os.path.join(target_desktop_folder, fname)
    if os.path.exists(src):
        shutil.copy2(src, dst)
        copied_files.append(dst)
        print(f"  [OK] Copied {fname}")
    else:
        print(f"  [ERROR] {fname} missing in {source_dir}")

print(f"Successfully copied {len(copied_files)} individual Excel files.")

# 2. Generate Combined Master Excel: ALL_PARTS_30_TO_51_MASTER.xlsx
print("\nGenerating combined master Excel file with all 35,123 voters...")
master_excel_path = os.path.join(target_desktop_folder, "ALL_PARTS_30_TO_51_MASTER.xlsx")

TALUK_MAP = {
    30: ("Molakalmuru", "Molakalmuru"),
    31: ("Molakalmuru", "Molakalmuru"),
    32: ("Challakere", "Challakere"),
    33: ("Challakere", "Challakere"),
    34: ("Challakere", "Challakere"),
    35: ("Challakere", "Challakere"),
    36: ("Chitradurga", "Chitradurga"),
    37: ("Chitradurga", "Chitradurga"),
    38: ("Chitradurga", "Chitradurga"),
    39: ("Chitradurga", "Chitradurga"),
    40: ("Hiriyur", "Hiriyur"),
    41: ("Hiriyur", "Hiriyur"),
    42: ("Hiriyur", "Hiriyur"),
    43: ("Hiriyur", "Hiriyur"),
    44: ("Hosadurga", "Hosadurga"),
    45: ("Hosadurga", "Hosadurga"),
    46: ("Hosadurga", "Hosadurga"),
    47: ("Hosadurga", "Hosadurga"),
    48: ("Holalkere", "Holalkere"),
    49: ("Holalkere", "Holalkere"),
    50: ("Holalkere", "Holalkere"),
    51: ("Holalkere", "Holalkere"),
}

def clean_name(val):
    if not val: return ""
    s = str(val).strip()
    if s.lower() in ('null', 'none', '—', '-', '.'): return ""
    s = re.sub(r'[\$#@!~%^*_+=\[\]{};:\"\'\<\>?/\\|]+', '', s).strip()
    s = re.sub(r'([a-z])([A-Z])', r'\1 \2', s)
    return re.sub(r'\s+', ' ', s).strip()

def clean_phone(val):
    if not val: return ""
    digits = re.sub(r'\D', '', str(val))
    if len(digits) == 10 and digits[0] in '6789':
        if len(set(digits)) > 1: return digits
    return ""

def clean_address(val):
    if not val: return ""
    s = str(val).strip()
    if s.lower() in ('null', 'none', '—', '-'): return ""
    s = s.replace('\\n', ', ').replace('\n', ', ').replace('\r', '')
    s = re.sub(r'[,;\s]+,', ',', s)
    s = re.sub(r',\s*,+', ', ', s)
    return re.sub(r'\s+', ' ', s).strip(' ,')

def clean_text(val):
    if not val: return ""
    s = str(val).strip()
    if s.lower() in ('null', 'none', '—', '-'): return ""
    return re.sub(r'\s+', ' ', s)

db_path = r'C:\Users\techb\Desktop\electrol-portal-NEW-main\electrol-portal-NEW-main\web\electors.db'
conn = sqlite3.connect(db_path)
c = conn.cursor()

c.execute("""
    SELECT 
        id,
        serial_number,
        epic_number,
        name,
        relative_name,
        sex,
        age,
        whatsapp_mob,
        address,
        part_number,
        polling_station_name,
        polling_address,
        qualification,
        occupation
    FROM electors
    WHERE CAST(part_number AS INTEGER) BETWEEN 30 AND 51
    ORDER BY 
        CAST(part_number AS INTEGER),
        CASE WHEN serial_number IS NOT NULL AND serial_number != '' THEN 0 ELSE 1 END,
        CAST(serial_number AS INTEGER),
        id
""")
all_rows = c.fetchall()
conn.close()

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Parts 30 to 51 Master"

headers = [
    "Sl No",
    "Roll Serial No",
    "EPIC Number",
    "Elector Name",
    "Relative Name",
    "Gender",
    "Age",
    "Mobile Number",
    "Address",
    "Part Number",
    "Polling Station Name",
    "Polling Address",
    "Taluk",
    "District",
    "Assembly Constituency",
    "Qualification",
    "Occupation"
]

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_side = Side(style='thin', color='D1D5DB')
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

regular_font = Font(name="Calibri", size=10)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

ws.append(headers)
ws.row_dimensions[1].height = 28

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = cell_border

phone_count = 0
center_cols = {1, 2, 3, 6, 7, 8, 10}

for idx, r in enumerate(all_rows, start=1):
    row_id, r_serial, r_epic, r_name, r_rel, r_sex, r_age, r_phone, r_addr, r_part, r_station, r_paddr, r_qual, r_occ = r
    
    p_num = int(r_part) if (r_part and str(r_part).isdigit()) else 30
    taluk_name, ac_name = TALUK_MAP.get(p_num, ("Chitradurga", "Chitradurga"))
    
    c_roll_serial = str(r_serial).strip() if (r_serial is not None and str(r_serial).strip() != '') else None
    c_epic = str(r_epic or '').strip().upper()
    c_name = clean_name(r_name)
    c_rel = clean_name(r_rel)
    c_sex = str(r_sex or '').strip().upper()
    if c_sex in ('M', 'MALE'): c_sex = 'M'
    elif c_sex in ('F', 'FEMALE'): c_sex = 'F'
    elif c_sex in ('T', 'THIRD', 'OTHER'): c_sex = 'T'
    
    c_age = int(r_age) if (r_age is not None and str(r_age).isdigit()) else None
    c_phone = clean_phone(r_phone)
    if c_phone:
        phone_count += 1
        phone_val = c_phone
    else:
        phone_val = None
        
    c_addr = clean_address(r_addr)
    c_station = clean_text(r_station)
    c_paddr = clean_text(r_paddr)
    c_qual = clean_text(r_qual)
    c_occ = clean_text(r_occ)
    
    row_data = [
        idx,
        c_roll_serial,
        c_epic,
        c_name,
        c_rel,
        c_sex,
        c_age,
        phone_val,
        c_addr,
        str(p_num),
        c_station,
        c_paddr,
        taluk_name,
        "Chitradurga",
        ac_name,
        c_qual,
        c_occ
    ]
    
    ws.append(row_data)
    current_row = idx + 1
    ws.row_dimensions[current_row].height = 20
    
    for c_idx in range(1, len(headers) + 1):
        c_cell = ws.cell(row=current_row, column=c_idx)
        c_cell.font = regular_font
        c_cell.border = cell_border
        c_cell.alignment = center_align if c_idx in center_cols else left_align
        if c_idx == 8 and phone_val:
            c_cell.number_format = '@'

ws.freeze_panes = "A2"
last_col_letter = get_column_letter(len(headers))
ws.auto_filter.ref = f"A1:{last_col_letter}{len(all_rows) + 1}"

min_widths = {
    1: 8, 2: 14, 3: 16, 4: 26, 5: 26, 6: 9, 7: 8, 8: 16, 9: 45,
    10: 12, 11: 35, 12: 40, 13: 16, 14: 15, 15: 22, 16: 18, 17: 20
}
for col_idx in range(1, len(headers) + 1):
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = min_widths.get(col_idx, 15)

wb.save(master_excel_path)
wb.close()
print(f"Master Excel file saved: {master_excel_path} ({len(all_rows):,} voters, {phone_count:,} phones)")

# 3. Create a Single ZIP Archive inside the folder
zip_path = os.path.join(target_desktop_folder, "Parts_30_to_51_All_Excel_Files.zip")
print(f"\nCreating ZIP archive of all files: {zip_path}...")
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
    for f in copied_files:
        zipf.write(f, os.path.basename(f))
    zipf.write(master_excel_path, os.path.basename(master_excel_path))

print(f"ZIP Archive created successfully ({os.path.getsize(zip_path):,} bytes)")

# 4. Mirror everything into the workspace exports directory as well
mirror_dir = r"C:\Users\techb\Desktop\electrol-portal-NEW-main\exports\Parts_30_to_51"
shutil.copy2(master_excel_path, os.path.join(mirror_dir, "ALL_PARTS_30_TO_51_MASTER.xlsx"))
shutil.copy2(zip_path, os.path.join(mirror_dir, "Parts_30_to_51_All_Excel_Files.zip"))

print(f"\nAll operations finished in {time.time() - start_time:.1f}s.")
print(f"FINAL FOLDER: {target_desktop_folder}")
