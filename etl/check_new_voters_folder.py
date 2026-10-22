import os
import glob
import openpyxl

target_epic = "YNT4252136"
target_num = "4252136"

dir_path = r"C:\Users\techb\Desktop\new voters"
files = glob.glob(os.path.join(dir_path, "*.xlsx"))
print(f"Scanning {len(files)} files in '{dir_path}'...")

found = False
for f in files:
    try:
        wb = openpyxl.load_workbook(f, read_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                r_str = " ".join([str(c) for c in row if c is not None])
                if target_num in r_str or target_epic in r_str:
                    print(f"FOUND in {os.path.basename(f)} (row {r_idx}):")
                    print("  ", row)
                    found = True
        wb.close()
    except Exception as e:
        print(f"Error {f}: {e}")

if not found:
    print(f"Not found in {dir_path}")
