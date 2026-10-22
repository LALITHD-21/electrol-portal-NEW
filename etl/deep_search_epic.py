import os
import glob
import openpyxl

target_epic = "YNT4252136"
target_number = "4252136"

print(f"=== DEEP SEARCH FOR '{target_epic}' or '{target_number}' ===\n")

# Check Excel files in root
root_excels = [
    'c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/MLC_Voters_ALL_Filtered_2026-10-05.xlsx',
    'c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/MLC_Voters_ALL_With_Phone_Numbers_2026.xlsx',
    'c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/new_members_added_list.xlsx',
    'c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/new_members_complete_list_2026.xlsx'
]

# Quick search in openpyxl / pandas
import pandas as pd
for ef in root_excels:
    if os.path.exists(ef):
        print(f"Checking {os.path.basename(ef)}...")
        try:
            # We can read chunk or read epic column
            wb = openpyxl.load_workbook(ef, read_only=True)
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    row_str = " ".join([str(v) for v in row if v is not None])
                    if target_number in row_str or target_epic in row_str:
                        print(f"  FOUND in {os.path.basename(ef)} (sheet {sheetname}, row {r_idx}):")
                        print("    ", row)
            wb.close()
        except Exception as e:
            print(f"  Error reading {ef}: {e}")

