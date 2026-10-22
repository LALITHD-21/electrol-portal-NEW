import openpyxl

wb = openpyxl.load_workbook('c:/Users/techb/Desktop/electrol-portal-NEW-main/electrol-portal-NEW-main/etl/test_out/Part_30_Voters.xlsx')
ws = wb.active
print(f"Sheet title: {ws.title}")
print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")

headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print(f"Headers ({len(headers)}):", headers)

print("\n--- First 3 data rows ---")
for r in range(2, 5):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Row {r-1}:", row_vals)

print("\n--- Check 5 rows with phone numbers ---")
count = 0
for r in range(2, ws.max_row + 1):
    phone = ws.cell(row=r, column=8).value
    if phone:
        print(f"Sl {ws.cell(row=r, column=1).value}: Name='{ws.cell(row=r, column=4).value}', Phone='{phone}' (type={type(phone).__name__}, len={len(str(phone))})")
        count += 1
        if count >= 5:
            break

print("\n--- Check 3 rows without phone numbers ---")
count = 0
for r in range(2, ws.max_row + 1):
    phone = ws.cell(row=r, column=8).value
    if not phone or phone == "":
        print(f"Sl {ws.cell(row=r, column=1).value}: Name='{ws.cell(row=r, column=4).value}', Phone='{phone}'")
        count += 1
        if count >= 3:
            break

wb.close()
